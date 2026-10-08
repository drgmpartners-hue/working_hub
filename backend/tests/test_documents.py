"""자료함 파서·처리 (기획 6장 '자료 읽기', P4-2).

파서는 이 파일 안에서 문서를 만들어 시험한다(pdf·docx·pptx·hwpx·xlsx·md·csv). 옛 형식:
- ppt: Apache POI 시험 파일(tests/fixtures/poi_basic_test.ppt, Apache License 2.0)
- hwp: 레코드 해석을 바이트로 직접 만들어 시험(실제 한글 파일은 HWP_SAMPLE 환경변수로 줄 때만)
처리 흐름(process)은 실제 PostgreSQL(PERM_PG_URL)에서 AI 호출을 가짜로 바꿔 시험한다.
"""
import io
import os
import struct
import uuid
import zipfile
from datetime import date
from pathlib import Path

import pytest

from app.services.company_report import doc_parser as dp

FIX = Path(__file__).parent / "fixtures"


def _png(w=400, h=300) -> bytes:
    from PIL import Image, ImageDraw

    im = Image.new("RGB", (w, h), (30, 80, 150))
    d = ImageDraw.Draw(im)
    for i in range(0, w, 13):
        d.line([i, 0, w - i, h], fill=(200, (i * 7) % 255, 60), width=2)
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


def _pdf(text_pages: list[str], with_image=True) -> bytes:
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    b = io.BytesIO()
    c = canvas.Canvas(b)
    for i, t in enumerate(text_pages):
        if t:
            c.setFont("Helvetica", 12)
            c.drawString(60, 780, t)
        if with_image and i == 0:
            c.drawImage(ImageReader(io.BytesIO(_png())), 60, 400, 300, 220)
        c.showPage()
    c.save()
    return b.getvalue()


def test_pdf_text_and_image():
    r = dp.parse("pdf", _pdf(["Revenue 12.0bn KRW in 2025", "Series B 15bn"]))
    assert r.status == "done" and r.page_count == 2 and not r.needs_ocr
    assert "Revenue 12.0bn" in r.text and "[2쪽]" in r.text
    assert len(r.images) == 1 and r.images[0].page == 1 and r.images[0].width == 400


def test_scanned_pdf_flagged():
    r = dp.parse("pdf", _pdf(["", ""], with_image=True))
    assert r.needs_ocr is True and r.status == "done"


def test_docx_body_table_image():
    import docx

    d = docx.Document()
    d.add_paragraph("대표이사 홍길동")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "한빛벤처스", "18%"
    d.add_paragraph("끝 문단")
    d.add_picture(io.BytesIO(_png()))
    b = io.BytesIO()
    d.save(b)
    r = dp.parse("docx", b.getvalue())
    assert r.text.splitlines() == ["대표이사 홍길동", "한빛벤처스 | 18%", "끝 문단"]
    assert len(r.images) == 1


def test_pptx_slides_notes_image():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "제품 라인업"
    s.shapes.add_picture(io.BytesIO(_png()), Inches(1), Inches(2))
    s.notes_slide.notes_text_frame.text = "유럽 인증 강조"
    b = io.BytesIO()
    prs.save(b)
    r = dp.parse("pptx", b.getvalue())
    assert "[슬라이드 1]" in r.text and "제품 라인업" in r.text and "(발표자 노트) 유럽 인증 강조" in r.text
    assert r.page_count == 1 and len(r.images) == 1 and r.images[0].page == 1


def test_hwpx_runs_and_bindata():
    sec = ('<?xml version="1.0"?><hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
           'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph"><hp:p><hp:run><hp:t>투자금 50억 원, </hp:t></hp:run>'
           '<hp:run><hp:t>3월 납입</hp:t></hp:run></hp:p></hs:sec>')
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr("Contents/section0.xml", sec)
        z.writestr("BinData/image1.png", _png())
        z.writestr("BinData/tiny.png", _png(20, 20))  # 아이콘 크기 — 버린다
    r = dp.parse("hwpx", b.getvalue())
    assert r.text == "투자금 50억 원, 3월 납입" and len(r.images) == 1


def test_hwp_record_parsing():
    # 문단 글자: '가' + 확장 제어(8글자 차지) + '나' + 문단 끝(13)
    payload = "가".encode("utf-16-le") + struct.pack("<H", 2) + b"\x00" * 14 + "나".encode("utf-16-le") + struct.pack("<H", 13)
    assert dp._hwp_para_text(payload) == "가나\n"
    header = (dp._HWPTAG_PARA_TEXT & 0x3FF) | (len(payload) << 20)
    recs = list(dp._hwp_records(struct.pack("<I", header) + payload))
    assert recs == [(dp._HWPTAG_PARA_TEXT, payload)]


@pytest.mark.skipif(not os.environ.get("HWP_SAMPLE"), reason="실제 .hwp 시험 파일 없음")
def test_real_hwp():
    r = dp.parse("hwp", Path(os.environ["HWP_SAMPLE"]).read_bytes())
    assert r.status == "done" and len(r.text) > 50


def test_ppt_records_skip_master_text():
    r = dp.parse("ppt", (FIX / "poi_basic_test.ppt").read_bytes())
    assert r.status == "done" and "This is a test title" in r.text
    assert "Click to edit Master" not in r.text and "___PPT10" not in r.text


def test_xlsx_csv_md_and_legacy():
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "손익"
    wb.active.append(["매출액", 12000000000])
    b = io.BytesIO()
    wb.save(b)
    assert "매출액 | 12000000000" in dp.parse("xlsx", b.getvalue()).text
    assert dp.parse("csv", "임직원수,85".encode("cp949")).text == "임직원수,85"
    assert dp.parse("md", "# 메모".encode()).text == "# 메모"
    r = dp.parse("doc", b"whatever")
    assert r.status == "unsupported" and ".docx" in r.error
    assert dp.parse("zip", b"x").status == "unsupported"
    assert dp.parse("pdf", b"not a pdf").status == "failed"


# --------------------------------------------------------------------------- 처리 흐름(실제 DB)

from tests.test_permissions import PG, env  # noqa: E402,F401


@pytest.mark.skipif(not PG, reason="PERM_PG_URL 없음(실제 PostgreSQL 필요)")
async def test_upload_registers_and_process(env, monkeypatch, tmp_path):  # noqa: F811
    import docx
    from sqlalchemy import select

    from app.models.company_report import CompanyDocument
    from app.services.company_report import documents

    monkeypatch.setenv("COMPANY_DB_ROOT", str(tmp_path))
    calls = []

    async def fake_memo(db, company, filename, text):
        calls.append((company, filename))
        return {"doc_type": "shareholders", "memo": "주주명부. 5번 투자유치에 쓸 만함", "facts": ["한빛벤처스 18%"],
                "has_personal_investment": True, "is_public": False, "doc_date": "2026-03-31"}

    monkeypatch.setattr(documents, "_memo", fake_memo)
    c, d, hdr = env["c"], env["d"], env["hdr"]
    ha, hb = hdr(d["A"]), hdr(d["B"])
    co = (await c.post("/company-report/companies", headers=ha, json={"name": f"자료함-{uuid.uuid4().hex[:6]}",
                                                                      "backfill_months": 0})).json()
    dd = docx.Document()
    dd.add_paragraph("주주명부: 한빛벤처스 18%, 김고객 1,000주")
    dd.add_picture(io.BytesIO(_png()))
    buf = io.BytesIO()
    dd.save(buf)
    r = await c.post("/company-report/db/files", headers=ha, data={"company_id": co["id"], "folder": "docs"},
                     files={"file": ("주주명부.docx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 201, r.text
    # 이미지 폴더 업로드는 자료함에 넣지 않는다
    r2 = await c.post("/company-report/db/files", headers=ha, data={"company_id": co["id"], "folder": "images"},
                      files={"file": ("사진.png", _png(), "image/png")})
    assert r2.status_code == 201
    async with env["Session"]() as db:
        docs = (await db.execute(select(CompanyDocument).where(CompanyDocument.company_id == co["id"]))).scalars().all()
        assert len(docs) == 1 and docs[0].filename == "주주명부.docx"
        out = await documents.process(db, docs[0].id)
        assert out.extract_status == "done" and "한빛벤처스 18%" in out.extracted_text
        assert out.doc_type == "shareholders" and out.has_personal_investment is True and out.ai_facts == ["한빛벤처스 18%"]
        assert len(out.extracted_images) == 1 and (tmp_path / out.extracted_images[0]["key"]).exists()
        assert out.doc_date == date(2026, 3, 31) and out.doc_date_source == "ai"  # AI 가 문서 속 기준일을 찾음(2026-10-08)
        doc_id = out.id
    lst = (await c.get(f"/company-report/companies/{co['id']}/documents", headers=ha)).json()
    assert [x["doc_type_label"] for x in lst["items"]] == ["주주명부·지분"] and lst["can_edit"] is True
    assert (await c.get(f"/company-report/documents/{doc_id}/images/0", headers=ha)).status_code == 200
    assert "한빛벤처스" in (await c.get(f"/company-report/documents/{doc_id}/text", headers=ha)).json()["text"]
    # 다른 매니저(목록에 없음)는 못 본다
    assert (await c.get(f"/company-report/companies/{co['id']}/documents", headers=hb)).status_code == 404
    assert (await c.get(f"/company-report/documents/{doc_id}/text", headers=hb)).status_code == 404
    assert (await c.patch(f"/company-report/documents/{doc_id}", headers=hb, json={"use_in_report": False})).status_code == 404
    r = await c.patch(f"/company-report/documents/{doc_id}", headers=ha, json={"use_in_report": False, "doc_type": "ir"})
    assert r.status_code == 200 and r.json()["use_in_report"] is False and r.json()["doc_type"] == "ir"
    assert (await c.patch(f"/company-report/documents/{doc_id}", headers=ha, json={"doc_type": "nope"})).status_code == 422
    # 자료 날짜: 담당자가 고치면 manual, 다시 읽어도 지킨다 / 미래 날짜는 안 됨 / 비우면 올린 날 기준
    r = await c.patch(f"/company-report/documents/{doc_id}", headers=ha, json={"doc_date": "2025-12-31"})
    assert r.json()["doc_date"] == "2025-12-31" and r.json()["doc_date_source"] == "manual"
    async with env["Session"]() as db:
        again = await documents.process(db, doc_id)
        assert again.doc_date == date(2025, 12, 31) and again.doc_date_source == "manual"
    assert (await c.patch(f"/company-report/documents/{doc_id}", headers=ha, json={"doc_date": "2099-01-01"})).status_code == 422
    r = await c.patch(f"/company-report/documents/{doc_id}", headers=ha, json={"doc_date": None})
    assert r.json()["doc_date"] is None and r.json()["doc_date_source"] is None
    assert documents.parse_doc_date("2026-02-30") is None and documents.parse_doc_date("") is None
    assert documents.parse_doc_date("2026-02-28 기준") == date(2026, 2, 28)
    # 검색: 본문 글자로 찾힌다
    res = (await c.get("/company-report/search", headers=ha, params={"q": "한빛벤처스"})).json()
    assert res["counts"].get("file", 0) >= 1, res["counts"]


@pytest.mark.skipif(not PG, reason="PERM_PG_URL 없음(실제 PostgreSQL 필요)")
async def test_scanned_pdf_uses_ai_reader(env, monkeypatch, tmp_path):  # noqa: F811
    from sqlalchemy import select

    from app.models.company_report import CompanyDocument
    from app.services.company_report import documents

    monkeypatch.setenv("COMPANY_DB_ROOT", str(tmp_path))

    async def fake_ocr(db, data, pages):
        return "[1쪽]\n감사보고서 매출액 120억 원"

    async def no_memo(*a, **k):
        return None

    monkeypatch.setattr(documents, "_ocr_pdf", fake_ocr)
    monkeypatch.setattr(documents, "_memo", no_memo)
    c, d, hdr = env["c"], env["d"], env["hdr"]
    ha = hdr(d["A"])
    co = (await c.post("/company-report/companies", headers=ha, json={"name": f"스캔-{uuid.uuid4().hex[:6]}",
                                                                      "backfill_months": 0})).json()
    r = await c.post("/company-report/db/files", headers=ha, data={"company_id": co["id"], "folder": "docs"},
                     files={"file": ("감사보고서.pdf", _pdf(["", ""]), "application/pdf")})
    assert r.status_code == 201
    async with env["Session"]() as db:
        doc = (await db.execute(select(CompanyDocument).where(CompanyDocument.company_id == co["id"]))).scalar_one()
        out = await documents.process(db, doc.id)
        assert out.extract_status == "done" and out.extract_method == "claude_pdf" and "120억" in out.extracted_text

    async def no_ocr(db, data, pages):
        return None

    monkeypatch.setattr(documents, "_ocr_pdf", no_ocr)  # 키가 없을 때
    async with env["Session"]() as db:
        out = await documents.process(db, doc.id)
        assert out.extract_status == "failed" and "스캔본" in out.extract_error
