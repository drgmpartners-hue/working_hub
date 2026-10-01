"""자료함 파서 — 문서에서 텍스트·이미지를 꺼낸다 (기획 6장 '자료 읽기', P4-2).

지원: pdf · docx · md · txt · pptx · hwpx · hwp · ppt · xlsx · csv
- LibreOffice 없이 순수 파이썬으로 처리한다(P4-3 결정). hwp·ppt(옛 형식)는 본문 글자 위주로 읽고,
  표·그림 배치가 중요한 문서는 hwpx·pptx·pdf 로 저장해 올리도록 화면에서 안내한다.
- doc·xls(옛 워드·엑셀)는 지원하지 않는다(unsupported) — docx·xlsx 로 저장해 다시 올린다.
- 스캔본 PDF(글자가 거의 없음)는 Claude 에 PDF 를 그대로 보내 읽는다(document_ai.py, 키가 있을 때).

반환값의 images 는 아직 저장 전 원본 바이트다. 너무 작은 그림(아이콘·로고 조각)은 버린다.
"""
from __future__ import annotations

import io
import logging
import re
import struct
import zipfile
import zlib
from dataclasses import dataclass, field
from typing import Optional
from xml.etree import ElementTree

logger = logging.getLogger(__name__)

MAX_TEXT = 400_000          # 저장하는 본문 최대 글자 수
MAX_PAGES = 400
MAX_IMAGES = 40
MIN_IMAGE_SIDE = 160        # 가로·세로 둘 다 이보다 작으면 버린다
MIN_IMAGE_BYTES = 4_000
SUPPORTED = {"pdf", "docx", "md", "txt", "pptx", "hwpx", "hwp", "ppt", "xlsx", "csv"}
LEGACY_HINT = {
    "doc": "옛 워드(.doc)는 읽지 못합니다. 워드에서 .docx 로 저장해 다시 올려 주세요.",
    "xls": "옛 엑셀(.xls)은 읽지 못합니다. 엑셀에서 .xlsx 로 저장해 다시 올려 주세요.",
}


@dataclass
class ParsedImage:
    data: bytes
    ext: str                      # png/jpg/gif/webp/bmp…
    page: Optional[int] = None    # 페이지·슬라이드 번호(1부터)
    name: Optional[str] = None
    width: Optional[int] = None
    height: Optional[int] = None


@dataclass
class ParseResult:
    text: str = ""
    page_count: Optional[int] = None
    images: list[ParsedImage] = field(default_factory=list)
    method: str = ""
    status: str = "done"          # done / failed / unsupported
    error: Optional[str] = None
    needs_ocr: bool = False       # 글자가 거의 없는 PDF(스캔본) — AI 로 다시 읽을 후보


# --------------------------------------------------------------------------- 공통

def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()[:MAX_TEXT]


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp949", "euc-kr"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


_IMG_EXT = {"png", "jpg", "jpeg", "gif", "bmp", "webp", "tif", "tiff", "emf", "wmf"}


def _image_from_bytes(data: bytes, ext: str, page: Optional[int] = None, name: Optional[str] = None) -> Optional[ParsedImage]:
    """크기 확인 + 웹에서 못 여는 형식(bmp·tif)은 PNG 로 바꾼다. emf·wmf(벡터)는 버린다."""
    ext = (ext or "").lower().lstrip(".")
    if ext == "jpeg":
        ext = "jpg"
    if ext in ("emf", "wmf") or len(data) < MIN_IMAGE_BYTES:
        return None
    try:
        from PIL import Image

        im = Image.open(io.BytesIO(data))
        w, h = im.size
        if w < MIN_IMAGE_SIDE and h < MIN_IMAGE_SIDE:
            return None
        if w < 60 or h < 60:  # 가는 선·띠
            return None
        if ext not in ("png", "jpg", "gif", "webp"):
            buf = io.BytesIO()
            im.convert("RGBA" if im.mode in ("RGBA", "LA", "P") else "RGB").save(buf, "PNG")
            data, ext = buf.getvalue(), "png"
        return ParsedImage(data=data, ext=ext, page=page, name=name, width=w, height=h)
    except Exception:
        return None


def _add_image(out: list[ParsedImage], img: Optional[ParsedImage], seen: set) -> None:
    if img is None or len(out) >= MAX_IMAGES:
        return
    key = (len(img.data), img.data[:64])
    if key in seen:  # 같은 그림이 여러 장에 반복(로고 등)
        return
    seen.add(key)
    out.append(img)


def _zip_images(z: zipfile.ZipFile, prefix: str) -> list[ParsedImage]:
    out: list[ParsedImage] = []
    seen: set = set()
    for n in sorted(z.namelist()):
        if not n.startswith(prefix):
            continue
        ext = n.rsplit(".", 1)[-1].lower() if "." in n else ""
        if ext not in _IMG_EXT:
            continue
        _add_image(out, _image_from_bytes(z.read(n), ext, name=n.rsplit("/", 1)[-1]), seen)
    return out


# --------------------------------------------------------------------------- 형식별

def parse_pdf(data: bytes) -> ParseResult:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            return ParseResult(status="failed", error="암호가 걸린 PDF 라 읽지 못했습니다. 암호를 풀어 다시 올려 주세요.", method="pypdf")
    n = len(reader.pages)
    parts: list[str] = []
    images: list[ParsedImage] = []
    seen: set = set()
    for i, page in enumerate(reader.pages[:MAX_PAGES], start=1):
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        if t.strip():
            parts.append(f"[{i}쪽]\n{t.strip()}")
        if len(images) < MAX_IMAGES:
            try:
                for im in page.images:
                    ext = (im.name.rsplit(".", 1)[-1] if "." in im.name else "png").lower()
                    _add_image(images, _image_from_bytes(im.data, ext, page=i, name=im.name), seen)
            except Exception:
                pass
    text = _clean("\n\n".join(parts))
    # 글자 있는 쪽이 40% 미만이거나 쪽당 평균 글자가 아주 적으면 스캔본으로 본다(IR 장표 PDF 는 글자가 적어도 괜찮음)
    pages = min(n, MAX_PAGES) or 1
    chars = len(re.sub(r"\s", "", text))
    needs_ocr = n > 0 and (len(parts) / pages < 0.4 or chars / pages < 15)
    return ParseResult(text=text, page_count=n, images=images, method="pypdf", needs_ocr=needs_ocr)


def parse_docx(data: bytes) -> ParseResult:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts: list[str] = []
    # 본문과 표를 문서 순서대로
    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = "".join(x.text or "" for x in child.iter() if x.tag.endswith("}t"))
            if t.strip():
                parts.append(t.strip())
        elif tag == "tbl":
            for tr in child.iter():
                if not tr.tag.endswith("}tr"):
                    continue
                cells = []
                for tc in tr:
                    if tc.tag.endswith("}tc"):
                        cells.append(" ".join("".join(x.text or "" for x in p.iter() if x.tag.endswith("}t"))
                                              for p in tc.iter() if p.tag.endswith("}p")).strip())
                if any(cells):
                    parts.append(" | ".join(cells))
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        images = _zip_images(z, "word/media/")
    return ParseResult(text=_clean("\n".join(parts)), images=images, method="python-docx")


def parse_pptx(data: bytes) -> ParseResult:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    images: list[ParsedImage] = []
    seen: set = set()
    for i, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            lines.extend(_pptx_shape_text(shape))
            if len(images) < MAX_IMAGES and getattr(shape, "shape_type", None) == 13:  # PICTURE
                try:
                    _add_image(images, _image_from_bytes(shape.image.blob, shape.image.ext, page=i), seen)
                except Exception:
                    pass
        notes = ""
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip() if slide.notes_slide.notes_text_frame else ""
        block = f"[슬라이드 {i}]\n" + "\n".join(x for x in lines if x.strip())
        if notes:
            block += f"\n(발표자 노트) {notes}"
        parts.append(block)
    return ParseResult(text=_clean("\n\n".join(parts)), page_count=len(prs.slides), images=images, method="python-pptx")


def _pptx_shape_text(shape) -> list[str]:
    out: list[str] = []
    if getattr(shape, "has_text_frame", False) and shape.text_frame:
        out.append(shape.text_frame.text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            out.append(" | ".join(c.text for c in row.cells))
    if getattr(shape, "shape_type", None) == 6:  # GROUP
        for s in shape.shapes:
            out.extend(_pptx_shape_text(s))
    return out


def parse_hwpx(data: bytes) -> ParseResult:
    """한글 2014+ (.hwpx): zip 안 Contents/section*.xml 의 글자, BinData 의 그림."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        sections = sorted((n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n)),
                          key=lambda n: int(re.findall(r"\d+", n)[-1]))
        parts: list[str] = []
        for n in sections:
            try:
                root = ElementTree.fromstring(z.read(n))
            except ElementTree.ParseError:
                continue
            for p in root.iter():
                if p.tag.rsplit("}", 1)[-1] != "p":
                    continue
                # 문단 바로 아래 run 의 글자만(표 안 문단은 그 문단에서 따로 잡힌다)
                t = "".join((x.text or "") for r in p if r.tag.rsplit("}", 1)[-1] == "run"
                            for x in r if x.tag.rsplit("}", 1)[-1] == "t")
                if t.strip():
                    parts.append(t.strip())
        images = _zip_images(z, "BinData/")
    return ParseResult(text=_clean("\n".join(parts)), page_count=len(sections) or None, images=images, method="hwpx-xml")


# HWP 5.0 레코드 태그
_HWPTAG_BEGIN = 0x10
_HWPTAG_PARA_TEXT = _HWPTAG_BEGIN + 51
# 8글자(16바이트)를 차지하는 제어 문자(확장·인라인)
_HWP_WIDE_CTRL = {1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23}


def _hwp_para_text(payload: bytes) -> str:
    out: list[str] = []
    i, n = 0, len(payload) - 1
    while i < n:
        code = payload[i] | (payload[i + 1] << 8)
        if code in _HWP_WIDE_CTRL:
            i += 16
            continue
        if code in (10, 13):
            out.append("\n")
        elif code == 9:
            out.append("\t")
        elif code >= 32:
            out.append(chr(code))
        i += 2
    return "".join(out)


def _hwp_records(buf: bytes):
    i, n = 0, len(buf)
    while i + 4 <= n:
        (h,) = struct.unpack_from("<I", buf, i)
        i += 4
        tag, size = h & 0x3FF, (h >> 20) & 0xFFF
        if size == 0xFFF:
            if i + 4 > n:
                break
            (size,) = struct.unpack_from("<I", buf, i)
            i += 4
        yield tag, buf[i:i + size]
        i += size


def parse_hwp(data: bytes) -> ParseResult:
    """한글 97~2010 (.hwp, HWP 5.0): OLE 안 BodyText/Section* 를 풀어 문단 글자를 읽는다."""
    import olefile

    if not olefile.isOleFile(io.BytesIO(data)):
        return ParseResult(status="failed", error="한글(.hwp) 파일 형식이 아닙니다.", method="hwp5")
    ole = olefile.OleFileIO(io.BytesIO(data))
    try:
        header = ole.openstream("FileHeader").read()
        flags = struct.unpack_from("<I", header, 36)[0] if len(header) >= 40 else 0
        compressed, encrypted = bool(flags & 1), bool(flags & 2)
        if encrypted:
            return ParseResult(status="failed", error="암호가 걸린 한글 문서라 읽지 못했습니다.", method="hwp5")
        sections = sorted((e for e in ole.listdir() if len(e) == 2 and e[0] == "BodyText" and e[1].startswith("Section")),
                          key=lambda e: int(re.findall(r"\d+", e[1])[0]))
        parts: list[str] = []
        for e in sections:
            raw = ole.openstream(e).read()
            if compressed:
                try:
                    raw = zlib.decompress(raw, -15)
                except zlib.error:
                    continue
            for tag, payload in _hwp_records(raw):
                if tag == _HWPTAG_PARA_TEXT:
                    t = _hwp_para_text(payload).strip()
                    if t:
                        parts.append(t)
        text = "\n".join(parts)
        if not text.strip() and ole.exists("PrvText"):  # 미리보기 글자(앞부분만)라도
            text = ole.openstream("PrvText").read().decode("utf-16-le", errors="ignore")
        images: list[ParsedImage] = []
        seen: set = set()
        for e in ole.listdir():
            if len(e) == 2 and e[0] == "BinData":
                raw = ole.openstream(e).read()
                if compressed:
                    try:
                        raw = zlib.decompress(raw, -15)
                    except zlib.error:
                        pass
                ext = e[1].rsplit(".", 1)[-1].lower() if "." in e[1] else ""
                _add_image(images, _image_from_bytes(raw, ext, name=e[1]), seen)
        return ParseResult(text=_clean(text), page_count=len(sections) or None, images=images, method="hwp5")
    finally:
        ole.close()


# PowerPoint 97~2003 레코드
_PPT_TEXT_CHARS = 0x0FA0     # UTF-16LE
_PPT_TEXT_BYTES = 0x0FA8     # 1바이트 문자
_PPT_SLIDE = 0x03EE
_PPT_CSTRING = 0x0FBA


def parse_ppt(data: bytes) -> ParseResult:
    """파워포인트 97~2003 (.ppt): 'PowerPoint Document' 스트림의 글자 레코드를 읽는다(그림은 꺼내지 않음)."""
    import olefile

    if not olefile.isOleFile(io.BytesIO(data)):
        return ParseResult(status="failed", error="파워포인트(.ppt) 파일 형식이 아닙니다.", method="ppt-records")
    ole = olefile.OleFileIO(io.BytesIO(data))
    try:
        if not ole.exists("PowerPoint Document"):
            return ParseResult(status="failed", error="파워포인트 본문을 찾지 못했습니다.", method="ppt-records")
        buf = ole.openstream("PowerPoint Document").read()
    finally:
        ole.close()
    parts: list[str] = []
    slides = 0

    def walk(start: int, end: int, depth: int = 0) -> None:
        nonlocal slides
        i = start
        while i + 8 <= end and depth < 20:
            ver_inst, rtype, rlen = struct.unpack_from("<HHI", buf, i)
            body_start, body_end = i + 8, min(i + 8 + rlen, end)
            if rtype == _PPT_SLIDE:
                slides += 1
            if (ver_inst & 0x0F) == 0x0F:  # 컨테이너
                walk(body_start, body_end, depth + 1)
            elif rtype == _PPT_TEXT_CHARS:
                parts.append(buf[body_start:body_end].decode("utf-16-le", errors="ignore"))
            elif rtype == _PPT_TEXT_BYTES:
                parts.append(buf[body_start:body_end].decode("cp1252", errors="ignore"))
            elif rtype == _PPT_CSTRING:
                s = buf[body_start:body_end].decode("utf-16-le", errors="ignore")
                if len(s) > 3:
                    parts.append(s)
            if rlen <= 0 and (ver_inst & 0x0F) != 0x0F:
                i = body_start
            else:
                i = body_start + rlen
            if i <= start:
                break

    walk(0, len(buf))
    text = "\n".join(p.replace("\r", "\n").strip() for p in parts if p and p.strip() and not _ppt_boilerplate(p))
    return ParseResult(text=_clean(text), page_count=slides or None, method="ppt-records")


_PPT_MASTER = re.compile(
    r"^(___PPT\d+|\*|Click to edit Master.*|Second level|Third level|Fourth level|Fifth level|Default Design|"
    r"마스터 제목 스타일 편집|마스터 텍스트 스타일(을 편집합니다)?|둘째 수준|셋째 수준|넷째 수준|다섯째 수준|Office 테마)$"
)


def _ppt_boilerplate(s: str) -> bool:
    """슬라이드 마스터의 자리 표시 글자(본문 아님)."""
    lines = [x.strip() for x in s.replace("\r", "\n").split("\n") if x.strip()]
    return bool(lines) and all(_PPT_MASTER.match(x) for x in lines)


def parse_xlsx(data: bytes) -> ParseResult:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for ws in wb.worksheets[:20]:
            rows = []
            for r in ws.iter_rows(values_only=True, max_row=2000):
                cells = ["" if v is None else str(v) for v in r]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                parts.append(f"[시트 {ws.title}]\n" + "\n".join(rows))
    finally:
        wb.close()
    return ParseResult(text=_clean("\n\n".join(parts)), page_count=len(wb.worksheets), method="openpyxl")


def parse_plain(data: bytes) -> ParseResult:
    return ParseResult(text=_clean(_decode(data)), method="text")


_PARSERS = {
    "pdf": parse_pdf, "docx": parse_docx, "pptx": parse_pptx, "hwpx": parse_hwpx, "hwp": parse_hwp, "ppt": parse_ppt,
    "xlsx": parse_xlsx, "md": parse_plain, "txt": parse_plain, "csv": parse_plain,
}


def parse(ext: str, data: bytes) -> ParseResult:
    """형식에 맞는 파서로 읽는다. 예외는 삼키고 failed 로 돌려준다(업로드는 그대로 둔다)."""
    ext = (ext or "").lower().lstrip(".")
    if ext in LEGACY_HINT:
        return ParseResult(status="unsupported", error=LEGACY_HINT[ext])
    fn = _PARSERS.get(ext)
    if fn is None:
        return ParseResult(status="unsupported", error=f".{ext or '?'} 형식은 자료함에서 읽지 않습니다.")
    try:
        r = fn(data)
    except Exception as e:  # 손상·특이 파일
        logger.info("자료 읽기 실패(%s): %s", ext, e)
        return ParseResult(status="failed", error=f"파일을 읽지 못했습니다: {type(e).__name__}", method=fn.__name__)
    if r.status == "done" and not r.text.strip() and not r.images and not r.needs_ocr:
        r.status, r.error = "failed", "글자를 찾지 못했습니다(빈 문서이거나 그림으로만 된 문서)."
    return r
