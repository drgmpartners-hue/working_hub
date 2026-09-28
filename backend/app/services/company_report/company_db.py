"""기업DB — 기업별 폴더·파일명 규칙·업로드·자동 파일·zip (기획 7-2, P2-5·6).

폴더(화면 이름): 01_기업정보 / 02_뉴스 / 03_자료 / 04_보고서 / 05_이미지, 그리고 '_포트폴리오 공통'
파일명 규칙: {기업명}_{문서종류}_{기간 또는 날짜}_{버전}.{확장자}  (버전은 2부터 v2로 붙인다)
"""
from __future__ import annotations

import io
import logging
import re
import uuid
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from typing import Optional
from xml.etree import ElementTree

from sqlalchemy import and_, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.company_report import CompanyFact, CompanyFile, CompanyFileDownload, CompanyFundingRound
from app.models.news_briefing import NewsArticle, NewsBriefing, PortfolioCompany
from app.services.company_report import search, storage
from app.services.company_report.timeutil import now_kst, today_kst

logger = logging.getLogger(__name__)

FOLDERS = {
    "info": "01_기업정보",
    "news": "02_뉴스",
    "docs": "03_자료",
    "reports": "04_보고서",
    "images": "05_이미지",
    "portfolio": "_포트폴리오 공통",
}
COMPANY_FOLDERS = ["info", "news", "docs", "reports", "images"]
PORTFOLIO_NAME = "_포트폴리오 공통"
DOC_EXT = {"pdf", "docx", "doc", "hwp", "hwpx", "ppt", "pptx", "md", "txt", "xlsx", "xls", "csv"}
IMG_EXT = {"png", "jpg", "jpeg", "gif", "webp"}
MAX_UPLOAD = 50 * 1024 * 1024
MIME = {
    "pdf": "application/pdf", "md": "text/markdown; charset=utf-8", "txt": "text/plain; charset=utf-8",
    "csv": "text/csv; charset=utf-8", "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "gif": "image/gif",
    "webp": "image/webp", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "hwp": "application/x-hwp", "hwpx": "application/hwp+zip", "zip": "application/zip",
}
_BAD = re.compile(r'[/\\:*?"<>|\x00-\x1f]')


# --------------------------------------------------------------------------- 이름 규칙

def sanitize(part: str, limit: int = 60) -> str:
    s = _BAD.sub("", part or "")
    s = re.sub(r"\s+", "", s).strip("._")
    return s[:limit]


def split_ext(filename: str) -> tuple[str, str]:
    name = filename.rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
    if "." in name:
        stem, ext = name.rsplit(".", 1)
        return stem, ext.lower()
    return name, ""


def display_name(company_name: Optional[str], kind: str, period: Optional[str], version: int, ext: str,
                 extra: Optional[str] = None) -> str:
    parts = [sanitize(company_name or PORTFOLIO_NAME.strip("_")) if company_name else None, sanitize(kind), sanitize(period or "")]
    if extra:
        parts.append(sanitize(extra, 40))
    if version and version > 1:
        parts.append(f"v{version}")
    base = "_".join(p for p in parts if p)
    return f"{base}.{ext}" if ext else base


def folder_names(companies: list[PortfolioCompany]) -> dict[str, str]:
    """기업 id → 화면 폴더 이름. 동명 기업은 사업자번호 끝 4자리(없으면 id 앞 4자리)를 붙인다."""
    by_name: dict[str, list[PortfolioCompany]] = defaultdict(list)
    for c in companies:
        by_name[sanitize(c.name, 80) or c.id[:8]].append(c)
    out = {}
    for name, cs in by_name.items():
        for c in cs:
            if len(cs) > 1:
                tail = re.sub(r"\D", "", c.biz_reg_no or "")[-4:] or c.id[:4]
                out[c.id] = f"{name}_{tail}"
            else:
                out[c.id] = name
    return out


# --------------------------------------------------------------------------- 본문 추출(검색용)

def _xml_text(data: bytes) -> str:
    try:
        root = ElementTree.fromstring(data)
    except ElementTree.ParseError:
        return ""
    return " ".join(t.strip() for t in root.itertext() if t and t.strip())


def extract_text(ext: str, data: bytes, limit: int = 20000) -> Optional[str]:
    """가벼운 텍스트 추출(검색 색인용). PDF·HWP 본문 추출은 자료함 파서(P4)에서 한다."""
    try:
        if ext in ("txt", "md", "csv"):
            for enc in ("utf-8", "cp949"):
                try:
                    return data.decode(enc)[:limit]
                except UnicodeDecodeError:
                    continue
            return None
        if ext in ("docx", "pptx", "hwpx", "xlsx"):
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = z.namelist()
                if ext == "docx":
                    parts = [n for n in names if n == "word/document.xml"]
                elif ext == "pptx":
                    parts = sorted(n for n in names if re.match(r"ppt/slides/slide\d+\.xml$", n))
                elif ext == "hwpx":
                    parts = sorted(n for n in names if re.match(r"Contents/section\d+\.xml$", n))
                else:
                    parts = [n for n in names if n == "xl/sharedStrings.xml"]
                text = " ".join(_xml_text(z.read(n)) for n in parts[:60])
                return text[:limit] or None
    except Exception as e:  # 손상 파일 등
        logger.info("본문 추출 실패(%s): %s", ext, e)
    return None


# --------------------------------------------------------------------------- 저장

async def _company_name(db: AsyncSession, company_id: Optional[str]) -> Optional[str]:
    if not company_id:
        return None
    c = await db.get(PortfolioCompany, company_id)
    return c.name if c else None


async def upload_file(db: AsyncSession, *, company_id: Optional[str], folder: str, filename: str, data: bytes,
                      user_id: Optional[str], doc_kind: Optional[str] = None, memo: Optional[str] = None,
                      period_label: Optional[str] = None) -> CompanyFile:
    if folder not in FOLDERS or (company_id and folder == "portfolio") or (not company_id and folder != "portfolio"):
        raise ValueError("폴더가 올바르지 않습니다.")
    if len(data) > MAX_UPLOAD:
        raise ValueError("파일은 50MB까지 올릴 수 있습니다.")
    stem, ext = split_ext(filename)
    allowed = IMG_EXT if folder == "images" else (DOC_EXT | IMG_EXT)
    if ext not in allowed:
        raise ValueError(f"올릴 수 없는 형식입니다(.{ext or '없음'}).")
    cname = await _company_name(db, company_id)
    kind = doc_kind or {"docs": "자료", "images": "이미지", "info": "기업정보", "news": "뉴스", "reports": "보고서", "portfolio": "자료"}[folder]
    period = period_label or today_kst().strftime("%Y%m%d")
    fid = str(uuid.uuid4())
    key = storage.make_key(company_id, folder, fid, ext)
    size = storage.save_bytes(key, data)
    f = CompanyFile(
        id=fid, company_id=company_id, folder=folder, original_name=filename[:300],
        display_name=display_name(cname, kind, period, 1, ext, extra=stem),
        storage_key=key, file_type=ext, mime=MIME.get(ext), size=size, period_label=period, doc_kind=kind,
        origin="upload", version=1, status="active", memo=memo, search_text=extract_text(ext, data), created_by=user_id,
    )
    db.add(f)
    await db.flush()
    await search.index_file(db, f, cname)
    return f


async def save_auto(db: AsyncSession, *, company_id: Optional[str], folder: str, doc_kind: str, period_label: Optional[str],
                    ext: str, data: bytes, auto_key: str, related_type: Optional[str] = None,
                    related_id: Optional[str] = None, search_text: Optional[str] = None) -> CompanyFile:
    """자동 파일 저장. 같은 auto_key가 있으면 내용만 바꾼다(보이는 이름·위치 유지)."""
    cname = await _company_name(db, company_id)
    f = (await db.execute(select(CompanyFile).where(CompanyFile.auto_key == auto_key, CompanyFile.status != "deleted"))).scalars().first()
    name = display_name(cname, doc_kind, period_label, 1, ext)
    if f:
        storage.save_bytes(f.storage_key, data)
        f.size, f.display_name, f.search_text, f.updated_at = len(data), name, search_text, now_kst()
    else:
        fid = str(uuid.uuid4())
        key = storage.make_key(company_id, folder, fid, ext)
        storage.save_bytes(key, data)
        f = CompanyFile(id=fid, company_id=company_id, folder=folder, display_name=name, storage_key=key, file_type=ext,
                        mime=MIME.get(ext), size=len(data), period_label=period_label, doc_kind=doc_kind, origin="auto",
                        auto_key=auto_key, related_type=related_type, related_id=related_id, version=1, status="active",
                        search_text=search_text)
        db.add(f)
        await db.flush()
    await search.index_file(db, f, cname)
    return f


# --------------------------------------------------------------------------- 자동 파일 생성

def _month_bounds(month: str) -> tuple[datetime, datetime]:
    y, m = (int(x) for x in month.split("-"))
    start = datetime(y, m, 1)
    end = datetime(y + (m == 12), m % 12 + 1, 1)
    return start, end


async def news_month_md(db: AsyncSession, company: PortfolioCompany, month: str) -> tuple[bytes, int]:
    start, end = _month_bounds(month)
    arts = (await db.execute(select(NewsArticle).where(
        NewsArticle.company_id == company.id, NewsArticle.is_hidden == False, NewsArticle.is_representative == True,  # noqa: E712
        NewsArticle.published_at >= start, NewsArticle.published_at < end,
    ).order_by(NewsArticle.published_at))).scalars().all()
    tag = {"positive": "호재", "neutral": "중립", "caution": "주의"}
    lines = [f"# {company.name} 뉴스 모음 — {month}", "", f"- 기사 {len(arts)}건 · 주의 {sum(1 for a in arts if a.tag == 'caution')}건",
             f"- 만든 시각: {now_kst().strftime('%Y-%m-%d %H:%M')} (매일 자동 갱신)", ""]
    cur = None
    for a in arts:
        d = a.published_at.strftime("%Y-%m-%d") if a.published_at else "날짜 미상"
        if d != cur:
            lines += ["", f"## {d}", ""]
            cur = d
        t = f"[{tag.get(a.tag or '', '-')}] " if a.tag else ""
        lines.append(f"- {t}**{a.title}** ({a.press or a.source})  ")
        if a.summary:
            lines.append(f"  {a.summary}  ")
        lines.append(f"  <{a.url}>")
    return "\n".join(lines).encode("utf-8"), len(arts)


def _xlsx(title: str, headers: list[str], rows: list[list], widths: list[int]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def facts_xlsx(db: AsyncSession, company: PortfolioCompany) -> tuple[bytes, int]:
    from app.services.company_report.facts import FACT_TYPES

    rows = (await db.execute(select(CompanyFact).where(CompanyFact.company_id == company.id, CompanyFact.status == "confirmed")
                             .order_by(CompanyFact.fact_date.desc().nullslast()))).scalars().all()
    data = [[f.fact_date.isoformat() if f.fact_date else "", FACT_TYPES.get(f.fact_type, f.fact_type), f.title,
             ", ".join(f"{k}: {v}" for k, v in (f.detail or {}).items() if not isinstance(v, (dict, list))),
             "\n".join(r.get("url", "") for r in (f.source_refs or [])[:3]), "직접 입력" if f.origin == "manual" else "기사"]
            for f in rows]
    return _xlsx("사실원장", ["날짜", "유형", "내용", "세부", "출처", "근거"], data, [12, 12, 50, 40, 60, 10]), len(rows)


async def funding_xlsx(db: AsyncSession, company: PortfolioCompany) -> tuple[bytes, int]:
    rows = (await db.execute(select(CompanyFundingRound).where(CompanyFundingRound.company_id == company.id,
                                                               CompanyFundingRound.status == "confirmed")
                             .order_by(CompanyFundingRound.round_date.nullslast()))).scalars().all()
    data = [[r.round_date.isoformat() if r.round_date else "", r.round_name or "", (r.amount / 1e8) if r.amount else "비공개",
             ", ".join(f"{i.get('name')}{'(리드)' if i.get('lead') else ''}" for i in (r.investors or [])),
             "예" if r.is_follow_on else "", "\n".join(s.get("url", "") for s in (r.source_refs or [])[:3])] for r in rows]
    return _xlsx("투자유치기록", ["날짜", "라운드", "금액(억 원)", "투자사", "후속", "출처"], data, [12, 12, 14, 50, 8, 60]), len(rows)


FONT_CANDIDATES = [
    # (보통, 굵게) — Docker(Railway)는 fonts-nanum 패키지, 로컬 Windows는 맑은 고딕
    ("/usr/share/fonts/truetype/nanum/NanumGothic.ttf", "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf"),
    ("C:/Windows/Fonts/malgun.ttf", "C:/Windows/Fonts/malgunbd.ttf"),
]


def _pdf_fonts() -> tuple[str, str]:
    """한글 TTF를 PDF에 포함(embed)한다. 없으면 CID 폰트(보는 프로그램에 따라 깨질 수 있음)."""
    import os

    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    try:
        pdfmetrics.getFont("CR-Regular")
        return "CR-Regular", "CR-Bold"
    except KeyError:
        pass
    env = os.environ.get("CR_PDF_FONT")
    cands = ([(env, os.environ.get("CR_PDF_FONT_BOLD", env))] if env else []) + FONT_CANDIDATES
    for reg, bold in cands:
        if reg and os.path.exists(reg):
            pdfmetrics.registerFont(TTFont("CR-Regular", reg))
            pdfmetrics.registerFont(TTFont("CR-Bold", bold if bold and os.path.exists(bold) else reg))
            return "CR-Regular", "CR-Bold"
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    logger.warning("한글 TTF 폰트를 찾지 못해 CID 폰트를 씁니다(Dockerfile에 fonts-nanum 필요).")
    try:
        pdfmetrics.getFont("HYGothic-Medium")
    except KeyError:
        pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))
    return "HYGothic-Medium", "HYGothic-Medium"


def _pdf(title: str, blocks: list[tuple[str, object]]) -> bytes:
    """간단한 A4 PDF. blocks: [("h", text) | ("p", text) | ("table", rows)]"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    font, bold = _pdf_fonts()
    h1 = ParagraphStyle("h1", fontName=bold, fontSize=18, leading=24, textColor=colors.HexColor("#1F3A5F"), spaceAfter=6)
    h2 = ParagraphStyle("h2", fontName=bold, fontSize=12, leading=18, textColor=colors.HexColor("#1F3A5F"), spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("b", fontName=font, fontSize=9.5, leading=15)
    small = ParagraphStyle("s", fontName=font, fontSize=8.5, leading=12)
    story = [Paragraph(escape(title), h1)]
    for kind, val in blocks:
        if kind == "h":
            story.append(Paragraph(escape(str(val)), h2))
        elif kind == "p":
            story.append(Paragraph(escape(str(val)).replace("\n", "<br/>"), body))
        elif kind == "table" and val:
            rows = [[Paragraph(escape(str(c)), small) for c in r] for r in val]  # type: ignore[union-attr]
            ncol = len(rows[0]) if rows else 0
            width = A4[0] - 36 * mm
            t = Table(rows, repeatRows=1, colWidths=[32 * mm, width - 32 * mm] if ncol == 2 else None)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF2F7")),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C7CDD6")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(t)
        story.append(Spacer(1, 3 * mm))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
                      title=title).build(story)
    return buf.getvalue()


async def company_card_pdf(db: AsyncSession, company: PortfolioCompany) -> bytes:
    from app.services.company_report.facts import FACT_TYPES

    facts = (await db.execute(select(CompanyFact).where(CompanyFact.company_id == company.id, CompanyFact.status == "confirmed")
                              .order_by(CompanyFact.fact_date.desc().nullslast()).limit(15))).scalars().all()
    rounds = (await db.execute(select(CompanyFundingRound).where(CompanyFundingRound.company_id == company.id,
                                                                 CompanyFundingRound.status == "confirmed")
                               .order_by(CompanyFundingRound.round_date.nullslast()))).scalars().all()
    since = now_kst() - timedelta(days=90)
    cnt = (await db.execute(select(func.count(), func.count().filter(NewsArticle.tag == "caution")).where(
        NewsArticle.company_id == company.id, NewsArticle.is_hidden == False, NewsArticle.is_representative == True,  # noqa: E712
        NewsArticle.published_at >= since))).one()
    info = [["항목", "내용"], ["대표자", company.ceo_name or "-"], ["업종", company.industry or "-"],
            ["설립일", company.established_at.isoformat() if company.established_at else "-"],
            ["상장", f"상장 {company.stock_code or ''}" if company.is_listed else "비상장"], ["주소", company.address or "-"],
            ["홈페이지", company.homepage or "-"], ["최근 90일 기사", f"{cnt[0]}건 (주의 {cnt[1]}건)"]]
    blocks: list[tuple[str, object]] = [("p", f"기준일 {today_kst().isoformat()} · Working Hub 기업 리포트 자동 생성"), ("h", "기본 정보"), ("table", info)]
    if rounds:
        blocks += [("h", "투자유치 기록(확정)"), ("table", [["날짜", "라운드", "금액", "투자사"]] + [
            [r.round_date.isoformat() if r.round_date else "-", r.round_name or "-",
             f"{r.amount / 1e8:,.1f}억 원" if r.amount else "비공개", ", ".join(i.get("name", "") for i in (r.investors or []))]
            for r in rounds])]
    if facts:
        blocks += [("h", "주요 사실(확정, 최근 15건)"), ("table", [["날짜", "유형", "내용"]] + [
            [f.fact_date.isoformat() if f.fact_date else "-", FACT_TYPES.get(f.fact_type, f.fact_type), f.title] for f in facts])]
    return _pdf(f"{company.name} 기업카드", blocks)


def briefing_pdf(b: NewsBriefing) -> bytes:
    info = b.basic_info or {}
    blocks: list[tuple[str, object]] = []
    w = info.get("weather") or {}
    blocks.append(("p", f"{b.briefing_date.isoformat()}({info.get('weekday', '')}) · 서울 {w.get('text') or '-'} · 기사 {b.article_count}건 · 주의 {b.caution_count}건"))
    mk = info.get("markets") or []
    if mk:
        blocks += [("h", "전일 증시"), ("table", [["지수", "시가", "종가", "변동", "%"]] + [
            [m.get("name"), m.get("open") or "-", m.get("close") or "-", m.get("change") or "-", m.get("change_pct") or "-"]
            if m.get("available") else [m.get("name"), "조회 실패", "", "", ""] for m in mk])]
    overall = (b.review_summary or {}).get("overall") or [{"text": t} for t in (b.overall_summary or "").split("\n") if t]
    blocks += [("h", "종합 브리핑"), ("p", "\n".join(f"· {x.get('text', '')}" for x in overall))]
    for c in b.company_summaries or []:
        blocks.append(("h", f"{c.get('name')} (기사 {c.get('article_count', 0)}건{', 주의 ' + str(c.get('caution_count')) if c.get('caution_count') else ''})"))
        lines = [c.get("one_liner") or ""]
        for a in [x for x in c.get("articles", []) if not x.get("more")][:3]:
            lines.append(f"- {a.get('title')} — {a.get('summary') or ''}\n  {a.get('url')}")
        blocks.append(("p", "\n".join(lines)))
    return _pdf(f"투자기업 데일리 브리핑 {b.briefing_date.isoformat()}", blocks)


async def refresh_auto_files(db: AsyncSession, company_id: str, months: Optional[list[str]] = None) -> dict:
    """기업 자동 파일 갱신: 월별 뉴스 모음(이번 달·지난달), 사실 원장·투자유치 엑셀, 기업카드 PDF."""
    company = await db.get(PortfolioCompany, company_id)
    if not company:
        return {}
    today = today_kst()
    if months is None:
        prev = (today.replace(day=1) - timedelta(days=1))
        months = [today.strftime("%Y-%m"), prev.strftime("%Y-%m")]
    made = []
    for m in months:
        data, n = await news_month_md(db, company, m)
        if n or m == today.strftime("%Y-%m"):
            await save_auto(db, company_id=company.id, folder="news", doc_kind="뉴스", period_label=m, ext="md", data=data,
                            auto_key=f"news_md:{company.id}:{m}", search_text=data.decode("utf-8")[:20000])
            made.append(f"news:{m}")
    data, n = await facts_xlsx(db, company)
    await save_auto(db, company_id=company.id, folder="info", doc_kind="사실원장", period_label=None, ext="xlsx", data=data,
                    auto_key=f"facts_xlsx:{company.id}")
    data, n2 = await funding_xlsx(db, company)
    await save_auto(db, company_id=company.id, folder="info", doc_kind="투자유치기록", period_label=None, ext="xlsx", data=data,
                    auto_key=f"funding_xlsx:{company.id}")
    await save_auto(db, company_id=company.id, folder="info", doc_kind="기업카드", period_label=today.strftime("%Y-%m"), ext="pdf",
                    data=await company_card_pdf(db, company), auto_key=f"card_pdf:{company.id}:{today.strftime('%Y-%m')}")
    made += ["facts_xlsx", "funding_xlsx", "card_pdf"]
    await db.commit()
    return {"company": company.name, "files": made, "facts": n, "rounds": n2}


async def save_briefing_pdf(db: AsyncSession, b: NewsBriefing) -> Optional[CompanyFile]:
    d = b.briefing_date
    f = await save_auto(db, company_id=None, folder="portfolio", doc_kind="데일리브리핑", period_label=d.isoformat(), ext="pdf",
                        data=briefing_pdf(b), auto_key=f"daily_pdf:{d.isoformat()}", related_type="daily", related_id=b.id,
                        search_text=(b.overall_summary or "")[:5000])
    return f


# --------------------------------------------------------------------------- 조회·zip

async def tree(db: AsyncSession) -> dict:
    companies = (await db.execute(select(PortfolioCompany).order_by(PortfolioCompany.name))).scalars().all()
    names = folder_names(list(companies))
    counts: dict[Optional[str], dict[str, int]] = defaultdict(dict)
    rows = (await db.execute(select(CompanyFile.company_id, CompanyFile.folder, func.count())
                             .where(CompanyFile.status != "deleted").group_by(CompanyFile.company_id, CompanyFile.folder))).all()
    for cid, folder, n in rows:
        counts[cid][folder] = n
    items = [{"id": c.id, "name": c.name, "folder_name": names[c.id], "is_active": c.is_active,
              "counts": counts.get(c.id, {}), "total": sum(counts.get(c.id, {}).values())} for c in companies]
    items.sort(key=lambda x: (not x["is_active"], x["folder_name"]))
    return {"folders": FOLDERS, "portfolio": {"name": PORTFOLIO_NAME, "total": sum(counts.get(None, {}).values())},
            "companies": items, "storage": storage.usage()}


def file_out(f: CompanyFile, company_name: Optional[str] = None, folder_name: Optional[str] = None) -> dict:
    return {"id": f.id, "company_id": f.company_id, "company_name": company_name, "folder": f.folder,
            "folder_label": FOLDERS.get(f.folder, f.folder), "path": f"{folder_name or PORTFOLIO_NAME}/{FOLDERS.get(f.folder, f.folder)}",
            "display_name": f.display_name, "original_name": f.original_name, "file_type": f.file_type, "size": f.size,
            "period_label": f.period_label, "doc_kind": f.doc_kind, "origin": f.origin, "version": f.version,
            "status": f.status, "memo": f.memo, "created_at": f.created_at.isoformat() if f.created_at else None,
            "updated_at": f.updated_at.isoformat() if f.updated_at else None, "related_type": f.related_type,
            "related_id": f.related_id}


async def list_files(db: AsyncSession, *, company_ids: Optional[list[str]] = None, portfolio: bool = False,
                     folder: Optional[str] = None, file_type: Optional[str] = None, origin: Optional[str] = None,
                     status: Optional[str] = None, date_from: Optional[date] = None, date_to: Optional[date] = None,
                     q: Optional[str] = None, include_inactive: bool = True, page: int = 1, size: int = 50) -> dict:
    cond = [CompanyFile.status != "deleted"] if not status else [CompanyFile.status == status]
    if portfolio:
        cond.append(CompanyFile.company_id.is_(None))
    elif company_ids:
        cond.append(CompanyFile.company_id.in_(company_ids))
    if folder:
        cond.append(CompanyFile.folder == folder)
    if file_type:
        types = {"image": list(IMG_EXT), "doc": ["docx", "doc", "hwp", "hwpx"], "slide": ["ppt", "pptx"],
                 "sheet": ["xlsx", "xls", "csv"]}.get(file_type, [file_type])
        cond.append(CompanyFile.file_type.in_(types))
    if origin:
        cond.append(CompanyFile.origin == origin)
    if date_from:
        cond.append(CompanyFile.created_at >= datetime.combine(date_from, datetime.min.time()))
    if date_to:
        cond.append(CompanyFile.created_at < datetime.combine(date_to + timedelta(days=1), datetime.min.time()))
    if not include_inactive:
        cond.append(or_(CompanyFile.company_id.is_(None), PortfolioCompany.is_active == True))  # noqa: E712
    if q and q.strip():
        for term in q.split():
            like = f"%{term}%"
            cond.append(or_(CompanyFile.display_name.ilike(like), CompanyFile.original_name.ilike(like),
                            CompanyFile.search_text.ilike(like), CompanyFile.memo.ilike(like), PortfolioCompany.name.ilike(like),
                            PortfolioCompany.name_en.ilike(like), func.cast(PortfolioCompany.aliases, _text()).ilike(like)))
    base = select(CompanyFile, PortfolioCompany).outerjoin(PortfolioCompany, PortfolioCompany.id == CompanyFile.company_id).where(and_(*cond))
    total = (await db.execute(select(func.count()).select_from(base.subquery()))).scalar_one()
    rows = (await db.execute(base.order_by(CompanyFile.updated_at.desc()).offset((page - 1) * size).limit(size))).all()
    companies = (await db.execute(select(PortfolioCompany))).scalars().all()
    names = folder_names(list(companies))
    return {"total": total, "page": page, "size": size,
            "items": [file_out(f, c.name if c else None, names.get(c.id) if c else None) for f, c in rows]}


def _text():
    from sqlalchemy import Text

    return Text()


async def log_download(db: AsyncSession, *, file_id: Optional[str], company_id: Optional[str], user_id: Optional[str], kind: str = "file") -> None:
    db.add(CompanyFileDownload(file_id=file_id, company_id=company_id, user_id=user_id, kind=kind))
    await db.commit()


async def zip_company(db: AsyncSession, company_id: Optional[str]) -> tuple[bytes, str]:
    """기업 폴더(또는 _포트폴리오 공통)를 같은 구조의 zip으로. 같은 이름은 (2) 번호를 붙인다."""
    companies = (await db.execute(select(PortfolioCompany))).scalars().all()
    names = folder_names(list(companies))
    top = names.get(company_id, PORTFOLIO_NAME) if company_id else PORTFOLIO_NAME
    cond = [CompanyFile.status != "deleted", CompanyFile.company_id == company_id if company_id else CompanyFile.company_id.is_(None)]
    files = (await db.execute(select(CompanyFile).where(*cond).order_by(CompanyFile.folder, CompanyFile.display_name))).scalars().all()
    buf = io.BytesIO()
    used: set[str] = set()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for f in files:
            if not storage.exists(f.storage_key):
                continue
            sub = FOLDERS.get(f.folder, f.folder) if company_id else (f.doc_kind or "기타")
            path = f"{top}/{sub}/{f.display_name}"
            n = 2
            while path in used:
                stem, ext = split_ext(f.display_name)
                path = f"{top}/{sub}/{stem}({n}).{ext}"
                n += 1
            used.add(path)
            z.writestr(path, storage.read_bytes(f.storage_key))
    return buf.getvalue(), f"{top}.zip"
